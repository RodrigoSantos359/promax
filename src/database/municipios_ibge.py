"""
Tabela auxiliar (Postgres) para resolver o código IBGE de um município a
partir do "Código Promax" — carregada uma vez a partir da planilha
`identi_municipio_ibge.xlsx` (colunas: UF, Codigo Promax, Descricao, IBGE).

Usada para enriquecer o campo "CODIGO MUNICIPIO" de arquivos de largura
fixa do PromaX (ex: rotina 01.20.12), cujo valor bruto é a sigla do UF
seguida do código do município SEM separador (ex: "MA00002404" — UF="MA",
código Promax=2404, obtido removendo os zeros à esquerda).
"""

import re
from pathlib import Path
from typing import Optional, Tuple

from sqlalchemy import text as sql_text

from config.settings import BASE_DIR, logger
from src.database.dynamic_tables import postgres_engine

TABLE_NAME = "municipios_ibge"
DEFAULT_XLSX_PATH = BASE_DIR / "identi_municipio_ibge.xlsx"

# "MA00002404" -> ("MA", "00002404") -> remove zeros à esquerda -> 2404
_CODIGO_MUNICIPIO_RE = re.compile(r"^([A-Za-z]{2})0*([1-9]\d*)$")


def ensure_table() -> None:
    """Cria a tabela auxiliar no Postgres, se ainda não existir."""
    with postgres_engine.begin() as conn:
        conn.execute(sql_text(
            f'CREATE TABLE IF NOT EXISTS "{TABLE_NAME}" ('
            "codigo_promax INTEGER PRIMARY KEY, "
            "uf TEXT, "
            "descricao TEXT, "
            "codigo_ibge INTEGER"
            ")"
        ))


def load_from_xlsx(xlsx_path: Optional[Path] = None) -> int:
    """
    Carrega (upsert, por `codigo_promax`) a planilha inteira na tabela
    auxiliar. Pode ser chamada de novo a qualquer momento para atualizar.

    Returns:
        int: quantidade de municípios processados.
    """
    import openpyxl

    path = Path(xlsx_path) if xlsx_path else DEFAULT_XLSX_PATH
    if not path.exists():
        raise FileNotFoundError(f"Planilha de municípios/IBGE não encontrada: {path}")

    ensure_table()

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    registros = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None or row[1] is None:
            continue
        uf, codigo_promax, descricao, codigo_ibge = row[0], row[1], row[2], row[3]
        registros.append({
            "codigo_promax": int(codigo_promax),
            "uf": str(uf).strip().upper(),
            "descricao": str(descricao).strip() if descricao else None,
            "codigo_ibge": int(codigo_ibge) if codigo_ibge is not None else None,
        })

    with postgres_engine.begin() as conn:
        for registro in registros:
            conn.execute(sql_text(
                f'INSERT INTO "{TABLE_NAME}" (codigo_promax, uf, descricao, codigo_ibge) '
                "VALUES (:codigo_promax, :uf, :descricao, :codigo_ibge) "
                "ON CONFLICT (codigo_promax) DO UPDATE SET "
                "uf = EXCLUDED.uf, descricao = EXCLUDED.descricao, codigo_ibge = EXCLUDED.codigo_ibge"
            ), registro)

    logger.info(f"[municipios_ibge] {len(registros)} município(s) carregado(s)/atualizado(s) de '{path.name}'")
    return len(registros)


def parse_codigo_municipio(raw: str) -> Optional[Tuple[str, int]]:
    """
    Extrai (uf, codigo_promax) do valor bruto do campo "CODIGO MUNICIPIO"
    do PromaX (ex: "MA00002404" -> ("MA", 2404)). Retorna None se o valor
    não bater com o padrão esperado (2 letras + dígitos).
    """
    if not raw:
        return None
    match = _CODIGO_MUNICIPIO_RE.match(raw.strip())
    if not match:
        return None
    uf, codigo_str = match.groups()
    return uf.upper(), int(codigo_str)


def lookup_codigo_ibge(codigo_promax: int) -> Optional[int]:
    """Busca o código IBGE correspondente a um código Promax de município."""
    with postgres_engine.begin() as conn:
        row = conn.execute(
            sql_text(f'SELECT codigo_ibge FROM "{TABLE_NAME}" WHERE codigo_promax = :cp'),
            {"cp": codigo_promax},
        ).fetchone()
    return row[0] if row else None


def load_lookup_map() -> dict:
    """
    Carrega toda a tabela auxiliar de uma vez em memória (codigo_promax ->
    codigo_ibge) — usado ao importar um arquivo inteiro, para não fazer uma
    consulta ao Postgres por linha.
    """
    with postgres_engine.begin() as conn:
        rows = conn.execute(sql_text(f'SELECT codigo_promax, codigo_ibge FROM "{TABLE_NAME}"')).fetchall()
    return {r[0]: r[1] for r in rows}


__all__ = [
    "TABLE_NAME",
    "ensure_table",
    "load_from_xlsx",
    "parse_codigo_municipio",
    "lookup_codigo_ibge",
    "load_lookup_map",
]
